"""
家計Pro - Excel レポート生成モジュール
担当: フロントエンド担当
Excelファイルへの出力、フォーマット、チャート生成
"""
import os
from datetime import datetime
from typing import List, Dict, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from config import (
    EXCEL_STYLES, CATEGORIES, get_all_expense_categories,
    get_category_group, get_category_info, INCOME_CATEGORIES,
)
from analyzer import (
    monthly_summary, multi_month_trend, budget_vs_actual,
    cost_effectiveness_analysis, anomaly_detection, yearly_summary,
)


THIN_BORDER = Border(
    left=Side(style="thin", color=EXCEL_STYLES["border_color"]),
    right=Side(style="thin", color=EXCEL_STYLES["border_color"]),
    top=Side(style="thin", color=EXCEL_STYLES["border_color"]),
    bottom=Side(style="thin", color=EXCEL_STYLES["border_color"]),
)

HEADER_FONT = Font(
    name="Arial", size=EXCEL_STYLES["header_font_size"],
    bold=True, color=EXCEL_STYLES["header_font"],
)
HEADER_FILL = PatternFill("solid", fgColor=EXCEL_STYLES["header_fill"])
BODY_FONT = Font(name="Arial", size=EXCEL_STYLES["body_font_size"])
TITLE_FONT = Font(
    name="Arial", size=EXCEL_STYLES["title_font_size"],
    bold=True, color=EXCEL_STYLES["header_fill"],
)
CENTER = Alignment(horizontal="center", vertical="center")
YEN_FMT = '#,##0"円"'
PCT_FMT = "0.0%"


def _style_header_row(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _style_data_cell(ws, row: int, col: int, fmt: str = None, fill_color: str = None):
    cell = ws.cell(row=row, column=col)
    cell.font = BODY_FONT
    cell.border = THIN_BORDER
    if fmt:
        cell.number_format = fmt
    if fill_color:
        cell.fill = PatternFill("solid", fgColor=fill_color)
    return cell


def _auto_width(ws, min_width: int = 10, max_width: int = 30):
    for col in ws.columns:
        max_len = min_width
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)) * 1.3 + 2, max_width))
        ws.column_dimensions[col_letter].width = max_len


def generate_full_report(year: int, month: int, output_path: str) -> str:
    wb = Workbook()

    _create_dashboard(wb, year, month)
    _create_expense_detail(wb, year, month)
    _create_income_detail(wb, year, month)
    _create_budget_sheet(wb, year, month)
    _create_trend_sheet(wb, year)
    _create_cost_effectiveness_sheet(wb, year, month)
    _create_anomaly_sheet(wb, year, month)

    wb.save(output_path)
    return output_path


def _create_dashboard(wb: Workbook, year: int, month: int):
    ws = wb.active
    ws.title = "ダッシュボード"
    ws.sheet_properties.tabColor = "1F4E79"

    ws.merge_cells("A1:G1")
    ws["A1"] = f"家計Pro ダッシュボード - {year}年{month}月"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 40

    summary = monthly_summary(year, month)

    ws["A3"] = "総収入"
    ws["B3"] = summary["total_income"]
    ws["C3"] = "総支出"
    ws["D3"] = summary["total_expense"]
    ws["E3"] = "収支"
    ws["F3"] = summary["balance"]

    for col in [1, 3, 5]:
        cell = ws.cell(row=3, column=col)
        cell.font = Font(name="Arial", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor=EXCEL_STYLES["subheader_fill"])
    for col in [2, 4, 6]:
        cell = ws.cell(row=3, column=col)
        cell.number_format = YEN_FMT
        cell.font = Font(name="Arial", size=12, bold=True)
        if col == 6:
            cell.font = Font(
                name="Arial", size=12, bold=True,
                color="006100" if summary["balance"] >= 0 else "9C0006",
            )

    ws["A4"] = "貯蓄率"
    ws["B4"] = summary["saving_rate"] / 100
    ws["B4"].number_format = PCT_FMT

    row = 6
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "カテゴリ別支出"
    ws[f"A{row}"].font = Font(name="Arial", size=13, bold=True, color=EXCEL_STYLES["header_fill"])
    row += 1

    headers = ["カテゴリ", "分類", "金額", "構成比"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(headers))
    row += 1

    cat_start_row = row
    for cat, info in sorted(
        summary["by_category"].items(),
        key=lambda x: x[1]["amount"],
        reverse=True,
    ):
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=info["group"])
        ws.cell(row=row, column=3, value=info["amount"])
        pct = info["amount"] / summary["total_expense"] if summary["total_expense"] > 0 else 0
        ws.cell(row=row, column=4, value=pct)

        group_color = {
            "固定費": EXCEL_STYLES["fixed_fill"],
            "変動費": EXCEL_STYLES["variable_fill"],
            "特別費": EXCEL_STYLES["special_fill"],
        }.get(info["group"], "FFFFFF")

        for col in range(1, 5):
            _style_data_cell(ws, row, col,
                             fmt=YEN_FMT if col == 3 else (PCT_FMT if col == 4 else None),
                             fill_color=group_color)
        row += 1

    if summary["by_category"]:
        pie = PieChart()
        pie.title = "支出構成比"
        pie.style = 10
        pie.width = 16
        pie.height = 12
        cats_ref = Reference(ws, min_col=1, min_row=cat_start_row, max_row=row - 1)
        vals_ref = Reference(ws, min_col=3, min_row=cat_start_row, max_row=row - 1)
        pie.add_data(vals_ref)
        pie.set_categories(cats_ref)
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showCatName = True
        ws.add_chart(pie, "F6")

    row += 1
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "分類別合計"
    ws[f"A{row}"].font = Font(name="Arial", size=13, bold=True, color=EXCEL_STYLES["header_fill"])
    row += 1

    group_headers = ["分類", "金額", "構成比"]
    for i, h in enumerate(group_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(group_headers))
    row += 1

    for grp_name, grp_total in summary["by_group"].items():
        ws.cell(row=row, column=1, value=grp_name)
        ws.cell(row=row, column=2, value=grp_total)
        pct = grp_total / summary["total_expense"] if summary["total_expense"] > 0 else 0
        ws.cell(row=row, column=3, value=pct)

        group_color = {
            "固定費": EXCEL_STYLES["fixed_fill"],
            "変動費": EXCEL_STYLES["variable_fill"],
            "特別費": EXCEL_STYLES["special_fill"],
        }.get(grp_name, "FFFFFF")
        for col in range(1, 4):
            _style_data_cell(ws, row, col,
                             fmt=YEN_FMT if col == 2 else (PCT_FMT if col == 3 else None),
                             fill_color=group_color)
        row += 1

    _auto_width(ws)


def _create_expense_detail(wb: Workbook, year: int, month: int):
    from data_manager import get_expenses
    ws = wb.create_sheet("支出明細")
    ws.sheet_properties.tabColor = "C00000"

    expenses = get_expenses(year, month)

    ws.merge_cells("A1:F1")
    ws["A1"] = f"支出明細 - {year}年{month}月"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    headers = ["日付", "カテゴリ", "分類", "金額", "メモ", "ID"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    _style_header_row(ws, 3, len(headers))

    row = 4
    if not expenses.empty:
        for _, exp in expenses.iterrows():
            ws.cell(row=row, column=1, value=exp["date"].strftime("%Y-%m-%d"))
            ws.cell(row=row, column=2, value=exp["category"])
            ws.cell(row=row, column=3, value=exp.get("group", ""))
            ws.cell(row=row, column=4, value=exp["amount"])
            ws.cell(row=row, column=5, value=exp.get("memo", ""))
            ws.cell(row=row, column=6, value=exp.get("id", ""))

            group_color = {
                "固定費": EXCEL_STYLES["fixed_fill"],
                "変動費": EXCEL_STYLES["variable_fill"],
                "特別費": EXCEL_STYLES["special_fill"],
            }.get(exp.get("group", ""), "FFFFFF")

            for col in range(1, 7):
                _style_data_cell(ws, row, col,
                                 fmt=YEN_FMT if col == 4 else None,
                                 fill_color=group_color)
            row += 1

    ws.cell(row=row + 1, column=3, value="合計")
    ws.cell(row=row + 1, column=3).font = Font(name="Arial", bold=True)
    ws.cell(row=row + 1, column=4).font = Font(name="Arial", bold=True)
    if row > 4:
        ws.cell(row=row + 1, column=4, value=f"=SUM(D4:D{row - 1})")
    else:
        ws.cell(row=row + 1, column=4, value=0)
    ws.cell(row=row + 1, column=4).number_format = YEN_FMT

    _auto_width(ws)
    ws.auto_filter.ref = f"A3:F{row - 1}" if row > 4 else "A3:F3"


def _create_income_detail(wb: Workbook, year: int, month: int):
    from data_manager import get_incomes
    ws = wb.create_sheet("収入明細")
    ws.sheet_properties.tabColor = "006100"

    incomes = get_incomes(year, month)

    ws.merge_cells("A1:E1")
    ws["A1"] = f"収入明細 - {year}年{month}月"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    headers = ["日付", "カテゴリ", "金額", "メモ", "ID"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    _style_header_row(ws, 3, len(headers))

    row = 4
    if not incomes.empty:
        for _, inc in incomes.iterrows():
            ws.cell(row=row, column=1, value=inc["date"].strftime("%Y-%m-%d"))
            ws.cell(row=row, column=2, value=inc["category"])
            ws.cell(row=row, column=3, value=inc["amount"])
            ws.cell(row=row, column=4, value=inc.get("memo", ""))
            ws.cell(row=row, column=5, value=inc.get("id", ""))
            for col in range(1, 6):
                _style_data_cell(ws, row, col,
                                 fmt=YEN_FMT if col == 3 else None,
                                 fill_color=EXCEL_STYLES["income_fill"])
            row += 1

    ws.cell(row=row + 1, column=2, value="合計")
    ws.cell(row=row + 1, column=2).font = Font(name="Arial", bold=True)
    if row > 4:
        ws.cell(row=row + 1, column=3, value=f"=SUM(C4:C{row - 1})")
    else:
        ws.cell(row=row + 1, column=3, value=0)
    ws.cell(row=row + 1, column=3).number_format = YEN_FMT
    ws.cell(row=row + 1, column=3).font = Font(name="Arial", bold=True)

    _auto_width(ws)


def _create_budget_sheet(wb: Workbook, year: int, month: int):
    ws = wb.create_sheet("予算vs実績")
    ws.sheet_properties.tabColor = "BF8F00"

    ws.merge_cells("A1:G1")
    ws["A1"] = f"予算 vs 実績 - {year}年{month}月"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    headers = ["カテゴリ", "分類", "予算", "実績", "差額", "達成率", "状態"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    _style_header_row(ws, 3, len(headers))

    bva = budget_vs_actual(year, month)
    row = 4
    for item in bva:
        if item["budget"] == 0 and item["actual"] == 0:
            continue

        ws.cell(row=row, column=1, value=item["category"])
        ws.cell(row=row, column=2, value=item["group"])
        ws.cell(row=row, column=3, value=item["budget"])
        ws.cell(row=row, column=4, value=item["actual"])
        ws.cell(row=row, column=5, value=item["diff"])
        ws.cell(row=row, column=6, value=item["rate"] / 100)
        ws.cell(row=row, column=7, value=item["status"])

        status_color = {
            "大幅超過": EXCEL_STYLES["warning_fill"],
            "超過": "FFEB9C",
            "適正": EXCEL_STYLES["good_fill"],
            "余裕あり": "D6E4F0",
        }.get(item["status"], "FFFFFF")

        for col in range(1, 8):
            fmt = YEN_FMT if col in (3, 4, 5) else (PCT_FMT if col == 6 else None)
            _style_data_cell(ws, row, col, fmt=fmt)
        ws.cell(row=row, column=7).fill = PatternFill("solid", fgColor=status_color)
        row += 1

    if row > 4:
        chart = BarChart()
        chart.type = "col"
        chart.title = "予算 vs 実績"
        chart.style = 10
        chart.width = 20
        chart.height = 12
        cats = Reference(ws, min_col=1, min_row=4, max_row=row - 1)
        budget_data = Reference(ws, min_col=3, min_row=3, max_row=row - 1)
        actual_data = Reference(ws, min_col=4, min_row=3, max_row=row - 1)
        chart.add_data(budget_data, titles_from_data=True)
        chart.add_data(actual_data, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        ws.add_chart(chart, "I3")

    _auto_width(ws)


def _create_trend_sheet(wb: Workbook, year: int):
    ws = wb.create_sheet("月次トレンド")
    ws.sheet_properties.tabColor = "4472C4"

    ws.merge_cells("A1:H1")
    ws["A1"] = f"月次トレンド分析 - {year}年"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    trend_data = multi_month_trend(year, 6)

    headers = ["年月", "総支出", "総収入", "収支", "貯蓄率"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    _style_header_row(ws, 3, len(headers))

    row = 4
    for t in trend_data:
        ws.cell(row=row, column=1, value=f"{t['year']}年{t['month']}月")
        ws.cell(row=row, column=2, value=t["total_expense"])
        ws.cell(row=row, column=3, value=t["total_income"])
        ws.cell(row=row, column=4, value=t["balance"])
        ws.cell(row=row, column=5, value=t["saving_rate"] / 100)

        for col in range(1, 6):
            fmt = YEN_FMT if col in (2, 3, 4) else (PCT_FMT if col == 5 else None)
            _style_data_cell(ws, row, col, fmt=fmt)

        if t["balance"] < 0:
            ws.cell(row=row, column=4).font = Font(
                name="Arial", size=EXCEL_STYLES["body_font_size"], color="9C0006"
            )
        row += 1

    if len(trend_data) > 1:
        line = LineChart()
        line.title = "収支推移"
        line.style = 10
        line.width = 20
        line.height = 12
        line.y_axis.numFmt = '#,##0'
        cats = Reference(ws, min_col=1, min_row=4, max_row=row - 1)
        expense_vals = Reference(ws, min_col=2, min_row=3, max_row=row - 1)
        income_vals = Reference(ws, min_col=3, min_row=3, max_row=row - 1)
        balance_vals = Reference(ws, min_col=4, min_row=3, max_row=row - 1)
        line.add_data(expense_vals, titles_from_data=True)
        line.add_data(income_vals, titles_from_data=True)
        line.add_data(balance_vals, titles_from_data=True)
        line.set_categories(cats)
        ws.add_chart(line, "G3")

    row += 2
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"] = "カテゴリ別月次推移"
    ws[f"A{row}"].font = Font(name="Arial", size=13, bold=True, color=EXCEL_STYLES["header_fill"])
    row += 1

    cat_headers = ["カテゴリ"] + [f"{t['year']}/{t['month']:02d}" for t in trend_data] + ["平均", "合計"]
    for i, h in enumerate(cat_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(cat_headers))
    cat_header_row = row
    row += 1

    cat_start_row = row
    for cat in get_all_expense_categories():
        amounts = []
        for t in trend_data:
            cat_data = t["by_category"].get(cat, {})
            amounts.append(cat_data.get("amount", 0))

        if sum(amounts) == 0:
            continue

        ws.cell(row=row, column=1, value=cat)
        for j, amt in enumerate(amounts):
            ws.cell(row=row, column=j + 2, value=amt)
            _style_data_cell(ws, row, j + 2, fmt=YEN_FMT)

        n = len(amounts)
        avg_col = n + 2
        sum_col = n + 3
        first_data_col = get_column_letter(2)
        last_data_col = get_column_letter(n + 1)
        ws.cell(row=row, column=avg_col, value=f"=AVERAGE({first_data_col}{row}:{last_data_col}{row})")
        ws.cell(row=row, column=sum_col, value=f"=SUM({first_data_col}{row}:{last_data_col}{row})")
        _style_data_cell(ws, row, avg_col, fmt=YEN_FMT)
        _style_data_cell(ws, row, sum_col, fmt=YEN_FMT)
        _style_data_cell(ws, row, 1)
        row += 1

    _auto_width(ws)


def _create_cost_effectiveness_sheet(wb: Workbook, year: int, month: int):
    ws = wb.create_sheet("費用対効果")
    ws.sheet_properties.tabColor = "ED7D31"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"費用対効果分析 - {year}年{month}月"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    ce = cost_effectiveness_analysis(year, month)

    ws["A3"] = "効率スコア"
    ws["B3"] = ce["efficiency_score"] / 100
    ws["B3"].number_format = PCT_FMT
    ws["C3"] = "節約可能額"
    ws["D3"] = ce["savings_potential"]
    ws["D3"].number_format = YEN_FMT

    for col in [1, 3]:
        ws.cell(row=3, column=col).font = Font(name="Arial", size=11, bold=True)
        ws.cell(row=3, column=col).fill = PatternFill("solid", fgColor=EXCEL_STYLES["subheader_fill"])
    for col in [2, 4]:
        ws.cell(row=3, column=col).font = Font(name="Arial", size=12, bold=True)

    row = 5
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"] = "無駄な支出（必要度が低く高額なもの）"
    ws[f"A{row}"].font = Font(name="Arial", size=13, bold=True, color="C00000")
    row += 1

    if ce["waste_items"]:
        headers = ["カテゴリ", "金額", "必要度(10点)", "円/必要度", "改善提案"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        _style_header_row(ws, row, len(headers))
        row += 1

        for item in ce["waste_items"]:
            ws.cell(row=row, column=1, value=item["category"])
            ws.cell(row=row, column=2, value=item["amount"])
            ws.cell(row=row, column=3, value=item["necessity_score"])
            ws.cell(row=row, column=4, value=item["cost_per_necessity"])
            ws.cell(row=row, column=5, value=item["suggestion"])

            for col in range(1, 6):
                fmt = YEN_FMT if col in (2, 4) else None
                _style_data_cell(ws, row, col, fmt=fmt, fill_color=EXCEL_STYLES["warning_fill"])
            row += 1
    else:
        ws.cell(row=row, column=1, value="無駄な支出は検出されませんでした。")
        ws.cell(row=row, column=1).font = Font(name="Arial", size=11, color="006100")
        row += 1

    row += 2
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"] = "全カテゴリ 費用対効果マトリックス"
    ws[f"A{row}"].font = Font(name="Arial", size=13, bold=True, color=EXCEL_STYLES["header_fill"])
    row += 1

    from config import NECESSITY_SCORES
    from data_manager import get_expenses
    expenses = get_expenses(year, month)

    matrix_headers = ["カテゴリ", "金額", "必要度", "コスパ指数", "評価"]
    for i, h in enumerate(matrix_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(matrix_headers))
    row += 1

    if not expenses.empty:
        for cat in get_all_expense_categories():
            cat_exp = expenses[expenses["category"] == cat]
            amt = int(cat_exp["amount"].sum()) if not cat_exp.empty else 0
            if amt == 0:
                continue

            necessity = NECESSITY_SCORES.get(cat, 5)
            cospa = necessity / (amt / 10000) if amt > 0 else 0

            if cospa >= 2:
                rating = "優秀"
                color = EXCEL_STYLES["good_fill"]
            elif cospa >= 1:
                rating = "適正"
                color = EXCEL_STYLES["neutral_fill"]
            else:
                rating = "要改善"
                color = EXCEL_STYLES["warning_fill"]

            ws.cell(row=row, column=1, value=cat)
            ws.cell(row=row, column=2, value=amt)
            ws.cell(row=row, column=3, value=necessity)
            ws.cell(row=row, column=4, value=round(cospa, 2))
            ws.cell(row=row, column=5, value=rating)

            for col in range(1, 6):
                fmt = YEN_FMT if col == 2 else None
                _style_data_cell(ws, row, col, fmt=fmt)
            ws.cell(row=row, column=5).fill = PatternFill("solid", fgColor=color)
            row += 1

    _auto_width(ws)


def _create_anomaly_sheet(wb: Workbook, year: int, month: int):
    ws = wb.create_sheet("異常値検出")
    ws.sheet_properties.tabColor = "9C0006"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"異常値検出レポート - {year}年{month}月"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    ws["A3"] = "過去6ヶ月の平均から大きく逸脱したカテゴリを検出します（1.5σ以上）"
    ws["A3"].font = Font(name="Arial", size=10, italic=True)

    anomalies = anomaly_detection(year, month)

    headers = ["カテゴリ", "今月", "平均", "標準偏差", "Z値", "方向"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=5, column=i, value=h)
    _style_header_row(ws, 5, len(headers))

    row = 6
    if anomalies:
        for a in anomalies:
            ws.cell(row=row, column=1, value=a["category"])
            ws.cell(row=row, column=2, value=a["current"])
            ws.cell(row=row, column=3, value=a["average"])
            ws.cell(row=row, column=4, value=a["std_dev"])
            ws.cell(row=row, column=5, value=a["z_score"])
            ws.cell(row=row, column=6, value=a["direction"])

            color = EXCEL_STYLES["warning_fill"] if a["direction"] == "増加" else EXCEL_STYLES["good_fill"]
            for col in range(1, 7):
                fmt = YEN_FMT if col in (2, 3, 4) else None
                _style_data_cell(ws, row, col, fmt=fmt, fill_color=color)
            row += 1
    else:
        ws.cell(row=row, column=1, value="異常値は検出されませんでした。")
        ws.cell(row=row, column=1).font = Font(name="Arial", size=11, color="006100")

    _auto_width(ws)
